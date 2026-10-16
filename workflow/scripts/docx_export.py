#!/usr/bin/env python3
"""Editable Word export for the doctoral-style manuscript.

The DOCX produced here is the single formal source: the PDF is exported from
this file by Word itself (``word_pipeline.py``), so fonts, page breaks,
numbering and figure placement in both files are identical.

Markdown conventions understood by the exporter (see
``workflow/references/doctoral-thesis-format.md``):

* ``# 标题``                      document title (cover page)
* ``## 摘要`` / ``## Abstract``    unnumbered front matter, Roman page numbers
* ``[[TOC]]``                      Word TOC field (updated by Word on export)
* ``## 第1章 …`` / ``### 1.1 …``  body headings, numbering written explicitly
* ``## 参考文献`` / ``## 附录A …``  back matter, each starts on a new page
* ``![图2-1 说明](path)``          figure, caption placed below
* ``表2-1 说明`` followed by a pipe table, optional ``注：…`` line after it
* ``$$ 公式 $$ (3-1)``             display equation with right-aligned number
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"

BODY_FONT = "宋体"
HEADING_FONT = "黑体"
LATIN_FONT = "Times New Roman"
BODY_PT = 12
LINE_PT = 20
TEXT_WIDTH_CM = 15.8
MAX_FIGURE_WIDTH_CM = 14.0
MAX_FIGURE_HEIGHT_CM = 9.5

FRONT_HEADINGS = ("摘要", "Abstract", "ABSTRACT", "目录")
BACK_PATTERN = re.compile(r"^(参考文献|附录|致谢)")
CHAPTER_PATTERN = re.compile(r"^(第\s*\d+\s*章|[一二三四五六七八九十]+、)")


def _set_fonts(font, east_asia: str, latin: str = LATIN_FONT) -> None:
    font.name = latin
    rpr = font._element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(f"{W_NS}eastAsia", east_asia)
    rfonts.set(f"{W_NS}ascii", latin)
    rfonts.set(f"{W_NS}hAnsi", latin)
    rfonts.set(f"{W_NS}cs", latin)


def _run(paragraph, text: str, *, font: str = BODY_FONT, size: float = BODY_PT,
         bold: bool | None = None, italic: bool | None = None, superscript: bool = False):
    from docx.shared import Pt

    run = paragraph.add_run(text)
    _set_fonts(run.font, font)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if superscript:
        run.font.superscript = True
    return run


def _style(style, east_asia: str, size: float, *, bold: bool = False) -> None:
    from docx.shared import Pt, RGBColor

    _set_fonts(style.font, east_asia)
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.italic = False
    style.font.color.rgb = RGBColor(0, 0, 0)
    p_pr = style._element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag.endswith("}pBdr"):
            p_pr.remove(child)


def _paragraph_format(paragraph, *, indent: bool = True, align: str = "justify",
                      line_pt: float = LINE_PT, before: float = 0, after: float = 0) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
    from docx.shared import Pt

    fmt = paragraph.paragraph_format
    fmt.line_spacing = Pt(line_pt)
    fmt.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.first_line_indent = Pt(BODY_PT * 2) if indent else Pt(0)
    fmt.widow_control = True
    paragraph.alignment = {
        "justify": WD_ALIGN_PARAGRAPH.JUSTIFY,
        "center": WD_ALIGN_PARAGRAPH.CENTER,
        "left": WD_ALIGN_PARAGRAPH.LEFT,
    }[align]


def _field(paragraph, instruction: str, placeholder: str = "1", size: float = 10.5) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    run = paragraph.add_run()
    _set_fonts(run.font, BODY_FONT)
    from docx.shared import Pt

    run.font.size = Pt(size)
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = f" {instruction} "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = placeholder
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instr, separate, text, end])


_INLINE = re.compile(r"(\*\*(.+?)\*\*|\$(.+?)\$|\[(\d+(?:[,，\-–]\d+)*)\])")


def _inline(paragraph, text: str, *, size: float = BODY_PT, font: str = BODY_FONT) -> None:
    """Bold, inline math (italic Latin) and GB/T 7714 superscript citations."""
    text = re.sub(r"\[([^\]]+)\]\((?!\s*\))[^)]*\)", r"\1", text).replace("`", "")
    cursor = 0
    for match in _INLINE.finditer(text):
        if match.start() > cursor:
            _run(paragraph, text[cursor:match.start()], font=font, size=size)
        if match.group(2) is not None:
            _run(paragraph, match.group(2), font=font, size=size, bold=True)
        elif match.group(3) is not None:
            _run(paragraph, match.group(3), font=font, size=size, italic=True)
        else:
            _run(paragraph, f"[{match.group(4)}]", font=font, size=size, superscript=True)
        cursor = match.end()
    if cursor < len(text):
        _run(paragraph, text[cursor:], font=font, size=size)


def _border(element, edges: dict[str, tuple[str, int]]) -> None:
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    borders = element.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        element.append(borders)
    for edge, (value, size) in edges.items():
        node = OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), value)
        node.set(qn("w:sz"), str(size))
        node.set(qn("w:space"), "0")
        node.set(qn("w:color"), "000000")
        borders.append(node)


def _parse_table(block: str) -> list[list[str]]:
    rows = []
    for line in block.splitlines():
        stripped = line.strip()
        if not stripped.startswith("|"):
            continue
        cells = [cell.strip() for cell in stripped.strip("|").split("|")]
        if cells and all(set(cell) <= set("-: ") for cell in cells):
            continue
        rows.append(cells)
    return rows


def _three_line_table(document, rows: list[list[str]]) -> None:
    """GB/T 7713-style three-line table: thick top/bottom, thin rule under header."""
    from docx.enum.table import WD_TABLE_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    columns = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=columns)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    tbl_pr = table._tbl.tblPr
    width = OxmlElement("w:tblW")
    width.set(qn("w:w"), "5000")
    width.set(qn("w:type"), "pct")
    tbl_pr.append(width)
    last = len(rows) - 1
    for r, row in enumerate(rows):
        tr_pr = table.rows[r]._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        tr_pr.append(cant)
        if r == 0:
            header = OxmlElement("w:tblHeader")
            tr_pr.append(header)
        for c in range(columns):
            cell = table.cell(r, c)
            cell.text = ""
            paragraph = cell.paragraphs[0]
            numeric = bool(re.fullmatch(r"[\d.,%±\-–~ ]+|NA|待实测|—", row[c] if c < len(row) else ""))
            _paragraph_format(paragraph, indent=False, align="center" if (r == 0 or numeric) else "left", line_pt=15)
            _inline(paragraph, row[c] if c < len(row) else "", size=10.5)
            if r == 0:
                for run in paragraph.runs:
                    run.bold = True
            tc_pr = cell._tc.get_or_add_tcPr()
            edges: dict[str, tuple[str, int]] = {"left": ("nil", 0), "right": ("nil", 0)}
            edges["top"] = ("single", 12) if r == 0 else ("nil", 0)
            edges["bottom"] = ("single", 12) if r == last else (("single", 6) if r == 0 else ("nil", 0))
            _border(tc_pr, edges)


def _caption(document, text: str, *, keep_with_next: bool) -> None:
    paragraph = document.add_paragraph(style="Caption")
    _paragraph_format(paragraph, indent=False, align="center", line_pt=16, before=3, after=3)
    paragraph.paragraph_format.keep_with_next = keep_with_next
    _inline(paragraph, text, size=10.5)


def _figure(document, alt: str, source: str, base_dirs: Iterable[Path]) -> None:
    from io import BytesIO

    from docx.shared import Cm, Pt
    from PIL import Image

    path = _resolve_asset(source, base_dirs)
    if path is None:
        raise FileNotFoundError(f"图形资产缺失：{source}")
    paragraph = document.add_paragraph()
    _paragraph_format(paragraph, indent=False, align="center", line_pt=LINE_PT, before=6)
    paragraph.paragraph_format.line_spacing = None  # single spacing so the picture is not clipped
    paragraph.paragraph_format.keep_with_next = True
    with Image.open(path) as source_image:
        image = source_image.convert("RGB")
        width_cm = MAX_FIGURE_WIDTH_CM
        if image.height / image.width * width_cm > MAX_FIGURE_HEIGHT_CM:
            width_cm = MAX_FIGURE_HEIGHT_CM * image.width / image.height
        # 300 dpi at the printed width: enough for print, keeps DOCX portable.
        target_px = int(width_cm / 2.54 * 300)
        if image.width > target_px:
            image = image.resize((target_px, round(image.height * target_px / image.width)), Image.Resampling.LANCZOS)
        stream = BytesIO()
        if path.suffix.lower() == ".png" and "chart" in path.stem:
            image.save(stream, format="PNG", optimize=True)
        else:
            image.save(stream, format="JPEG", quality=90, optimize=True)
        stream.seek(0)
    paragraph.add_run().add_picture(stream, width=Cm(width_cm))
    _caption(document, alt, keep_with_next=False)
    document.paragraphs[-1].paragraph_format.space_after = Pt(6)


def _resolve_asset(source: str, base_dirs: Iterable[Path]) -> Path | None:
    raw = Path(source.strip().strip("<>"))
    candidates = [raw] if raw.is_absolute() else [Path(d) / raw for d in base_dirs] + [Path.cwd() / raw]
    for candidate in candidates:
        if candidate.is_file():
            return candidate.resolve()
    return None


def _equation(document, body: str, number: str) -> None:
    from docx.enum.text import WD_TAB_ALIGNMENT
    from docx.shared import Cm

    paragraph = document.add_paragraph()
    _paragraph_format(paragraph, indent=False, align="left", before=3, after=3)
    stops = paragraph.paragraph_format.tab_stops
    stops.add_tab_stop(Cm(TEXT_WIDTH_CM / 2), WD_TAB_ALIGNMENT.CENTER)
    stops.add_tab_stop(Cm(TEXT_WIDTH_CM), WD_TAB_ALIGNMENT.RIGHT)
    _run(paragraph, "\t")
    _run(paragraph, body.strip(), font=LATIN_FONT, italic=False)
    _run(paragraph, f"\t{number}")


def _page_number_footer(section, fmt: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    section.footer.is_linked_to_previous = False
    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _field(footer, "PAGE")
    sect_pr = section._sectPr
    pg = sect_pr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sect_pr.append(pg)
    pg.set(qn("w:fmt"), fmt)
    pg.set(qn("w:start"), "1")


def _header(section, text: str) -> None:
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn

    section.header.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.text = ""
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _run(header, text, size=10.5)
    p_pr = header._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for key, value in (("val", "single"), ("sz", "6"), ("space", "1"), ("color", "000000")):
        bottom.set(qn(f"w:{key}"), value)
    borders.append(bottom)
    p_pr.append(borders)


def _page_setup(section) -> None:
    from docx.shared import Cm

    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(3.0), Cm(2.5)
    section.left_margin = section.right_margin = Cm(2.6)
    section.header_distance, section.footer_distance = Cm(2.0), Cm(1.5)


def split_blocks(markdown: str) -> list[str]:
    return [part.strip("\n") for part in re.split(r"\n\s*\n", markdown.strip()) if part.strip()]


def export_docx(markdown: str, output_path: Path, *, title: str = "", subject: str = "",
                base_dirs: Iterable[Path] = ()) -> dict:
    """Create an editable DOCX. Raises on missing assets instead of hiding them."""
    from docx import Document
    from docx.enum.section import WD_SECTION
    from docx.enum.text import WD_BREAK
    from docx.shared import Pt

    base_dirs = list(base_dirs)
    document = Document()
    _page_setup(document.sections[0])
    styles = document.styles
    _style(styles["Normal"], BODY_FONT, BODY_PT)
    _style(styles["Title"], HEADING_FONT, 22, bold=True)
    for name, size in (("Heading 1", 16), ("Heading 2", 14), ("Heading 3", 12)):
        _style(styles[name], HEADING_FONT, size, bold=False)
        fmt = styles[name].paragraph_format
        fmt.keep_with_next = True
        fmt.space_before = Pt({"Heading 1": 18, "Heading 2": 12, "Heading 3": 6}[name])
        fmt.space_after = Pt({"Heading 1": 18, "Heading 2": 6, "Heading 3": 3}[name])
    _style(styles["Caption"], BODY_FONT, 10.5)
    for level in (1, 2, 3):
        if f"TOC {level}" in [s.name for s in styles]:
            _style(styles[f"TOC {level}"], BODY_FONT, 12)

    blocks = split_blocks(markdown)
    title_text = title
    if blocks and blocks[0].startswith("# "):
        title_text = blocks.pop(0)[2:].strip()

    # Cover page (no header/number).
    cover = document.sections[0]
    for _ in range(6):
        document.add_paragraph()
    paragraph = document.add_paragraph(style="Title")
    _paragraph_format(paragraph, indent=False, align="center", line_pt=36)
    _run(paragraph, title_text, font=HEADING_FONT, size=22, bold=True)
    cover_lines = []
    while blocks and blocks[0].startswith("> "):
        cover_lines.extend(line[2:].strip() for line in blocks.pop(0).splitlines())
    for _ in range(8):
        document.add_paragraph()
    for line in cover_lines:
        p = document.add_paragraph()
        _paragraph_format(p, indent=False, align="center", line_pt=28)
        _inline(p, line, size=14)

    state = "front"
    section = document.add_section(WD_SECTION.NEW_PAGE)
    _header(section, title_text[:40])
    _page_number_footer(section, "upperRoman")
    cover.different_first_page_header_footer = False
    first_in_section = True
    pending_table_caption: str | None = None
    counts = {"figures": 0, "tables": 0, "equations": 0, "headings": 0}

    def new_page() -> None:
        nonlocal first_in_section
        if not first_in_section:
            document.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
            document.paragraphs[-1].paragraph_format.space_after = Pt(0)
        first_in_section = False

    for index, block in enumerate(blocks):
        stripped = block.strip()
        if stripped == "[[TOC]]":
            new_page()
            p = document.add_paragraph(style="Heading 1")
            p.alignment = 1
            _run(p, "目　　录", font=HEADING_FONT, size=16)
            toc = document.add_paragraph()
            _field(toc, 'TOC \\o "1-3" \\h \\z \\u', "（目录由 Word 更新域生成）", size=12)
            continue
        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading and "\n" not in stripped:
            level = len(heading.group(1)) - 1
            text = heading.group(2).strip()
            counts["headings"] += 1
            if level == 1:
                if state == "front" and CHAPTER_PATTERN.match(text):
                    state = "body"
                    section = document.add_section(WD_SECTION.NEW_PAGE)
                    _page_number_footer(section, "decimal")
                    first_in_section = True
                elif state == "body" and BACK_PATTERN.match(text):
                    state = "back"
                new_page()
            p = document.add_paragraph(style=f"Heading {level}")
            if level == 1:
                p.alignment = 1
            _run(p, text, font=HEADING_FONT, size={1: 16, 2: 14, 3: 12}[level])
            continue
        first_in_section = False
        image = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image:
            _figure(document, image.group(1), image.group(2), base_dirs)
            counts["figures"] += 1
            continue
        equation = re.fullmatch(r"\$\$(.+?)\$\$\s*(（?\(?\d+-\d+\)?）?)", stripped, re.S)
        if equation:
            _equation(document, equation.group(1), equation.group(2))
            counts["equations"] += 1
            continue
        lines = [line for line in stripped.splitlines() if line.strip()]
        if lines and re.match(r"^表\s*\d+-\d+", lines[0]) and len(lines) > 1 and lines[1].strip().startswith("|"):
            _caption(document, lines[0].strip(), keep_with_next=True)
            lines = lines[1:]
            stripped = "\n".join(lines)
        elif len(lines) == 1 and re.match(r"^表\s*\d+-\d+", lines[0]):
            pending_table_caption = lines[0]
            continue
        if all(line.strip().startswith("|") for line in lines):
            if pending_table_caption:
                _caption(document, pending_table_caption, keep_with_next=True)
                pending_table_caption = None
            rows = _parse_table(stripped)
            if rows:
                _three_line_table(document, rows)
                counts["tables"] += 1
            continue
        if lines and lines[0].startswith("注："):
            p = document.add_paragraph()
            _paragraph_format(p, indent=False, align="justify", line_pt=16, before=2, after=6)
            _inline(p, " ".join(line.strip() for line in lines), size=10.5)
            continue
        if all(line.startswith(("- ", "* ")) for line in lines):
            for line in lines:
                p = document.add_paragraph()
                _paragraph_format(p, indent=True)
                _inline(p, "· " + line[2:].strip())
            continue
        if all(re.match(r"^\d+[.、)]\s", line) for line in lines):
            for line in lines:
                p = document.add_paragraph()
                _paragraph_format(p, indent=True)
                _inline(p, line.strip())
            continue
        if all(line.startswith("> ") for line in lines):
            for line in lines:
                p = document.add_paragraph()
                _paragraph_format(p, indent=False)
                _inline(p, line[2:])
            continue
        if state == "back" and re.match(r"^\[\d+\]", lines[0]):
            for line in lines:
                p = document.add_paragraph()
                _paragraph_format(p, indent=False, line_pt=18)
                p.paragraph_format.left_indent = Pt(24)
                p.paragraph_format.first_line_indent = Pt(-24)
                _run(p, line.strip(), size=10.5)
            continue
        p = document.add_paragraph()
        _paragraph_format(p)
        _inline(p, "".join(line.strip() for line in lines))

    document.core_properties.title = title_text
    document.core_properties.subject = subject
    document.core_properties.comments = "Generated by autoartical docx_export (single formal source)"
    settings = document.settings.element
    update = settings.find(f"{W_NS}updateFields")
    if update is None:
        from docx.oxml import OxmlElement

        update = OxmlElement("w:updateFields")
        settings.append(update)
    update.set(f"{W_NS}val", "true")
    output_path.parent.mkdir(parents=True, exist_ok=True)
    document.save(output_path)
    return {"path": str(output_path), "format": "docx", "editable": True, **counts}
